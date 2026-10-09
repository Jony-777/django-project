# -*- coding: utf-8 -*-
"""Genera capturas y el Word del RA Unidad 1 con imágenes embebidas."""
from __future__ import annotations

import json
import os
import subprocess
import sys
import textwrap
import time
from pathlib import Path

import requests
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

BASE = Path(__file__).resolve().parent
CAPTURAS = BASE / "informe_capturas"
OUTPUT = BASE / "Resultado_Aprendizaje_Unidad1_con_capturas.docx"
FONT_PATH = Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts" / "consola.ttf"


def set_cell_shading(cell, fill_hex: str):
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill_hex)
    cell._tc.get_or_add_tcPr().append(shading)


def load_font(size: int) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    if FONT_PATH.exists():
        return ImageFont.truetype(str(FONT_PATH), size)
    return ImageFont.load_default()


def render_code_image(
    path: Path,
    *,
    title: str,
    start_line: int | None = None,
    end_line: int | None = None,
    highlight_lines: set[int] | None = None,
) -> Path:
    text = path.read_text(encoding="utf-8")
    lines = text.splitlines()
    if start_line is not None and end_line is not None:
        lines = lines[start_line - 1 : end_line]
        base_line = start_line
    else:
        base_line = 1

    font = load_font(16)
    line_h = 22
    pad_x, pad_y = 16, 48
    max_chars = max(len(line) for line in lines) if lines else 40
    width = min(1400, pad_x * 2 + max_chars * 10)
    height = pad_y + len(lines) * line_h + 16

    img = Image.new("RGB", (width, height), "#1e1e1e")
    draw = ImageDraw.Draw(img)
    draw.rectangle((0, 0, width, 36), fill="#2d2d30")
    draw.text((12, 8), title, fill="#cccccc", font=load_font(14))

    highlight_lines = highlight_lines or set()
    for i, line in enumerate(lines):
        y = pad_y + i * line_h
        ln = base_line + i
        if ln in highlight_lines:
            draw.rectangle((0, y - 2, width, y + line_h - 2), fill="#264f78")
        draw.text((pad_x, y), f"{ln:4}  {line}", fill="#d4d4d4", font=font)

    safe = "".join(c if c.isalnum() or c in "-_" else "_" for c in title)
    out = CAPTURAS / f"code_{path.stem}_{safe}.png"
    img.save(out)
    return out


def render_terminal_image(output_text: str, title: str = "PowerShell — python manage.py test products") -> Path:
    font = load_font(15)
    lines = output_text.splitlines()
    line_h = 20
    width, height = 920, 60 + len(lines) * line_h
    img = Image.new("RGB", (width, height), "#0c0c0c")
    draw = ImageDraw.Draw(img)
    draw.rectangle((0, 0, width, 32), fill="#1f1f1f")
    draw.text((10, 7), title, fill="#cccccc", font=load_font(13))
    for i, line in enumerate(lines):
        color = "#4ec9b0" if "OK" in line else "#dcdcdc"
        if "FAIL" in line or "Error" in line:
            color = "#f44747"
        draw.text((12, 40 + i * line_h), line, fill=color, font=font)
    out = CAPTURAS / "captura_terminal_tests.png"
    img.save(out)
    return out


def render_json_api_image(
    method: str,
    url: str,
    status: int,
    body: dict,
    headers_note: str,
) -> Path:
    font = load_font(14)
    json_str = json.dumps(body, indent=2, ensure_ascii=False)
    lines = [
        f"{method} {url}",
        f"Status: {status}",
        headers_note,
        "",
        json_str,
    ]
    line_h = 20
    width = 880
    height = 80 + len("\n".join(lines).splitlines()) * line_h
    img = Image.new("RGB", (width, height), "#f5f5f5")
    draw = ImageDraw.Draw(img)
    draw.rectangle((0, 0, width, 44), fill="#ff6c37")
    draw.text((12, 12), "Cliente HTTP (Postman / API)", fill="white", font=load_font(15))
    y = 52
    for block in lines:
        for sub in block.splitlines():
            draw.text((14, y), sub, fill="#333333", font=font)
            y += line_h
    parts = [p for p in url.rstrip("/").split("/") if p]
    product_id = parts[-2] if parts and parts[-1] == "precio-final" else parts[-1]
    out = CAPTURAS / f"captura_api_producto_{product_id}.png"
    img.save(out)
    return out


def render_diagram_image() -> Path:
    diagram = textwrap.dedent(
        """
        Flujo GET /api/v1/products/<id>/precio-final/  (JWT)

          Cliente (Postman)
                |
                v
          urls.py  (enrutamiento)
                |
                v
          views.py  precio_final  [Facade]
                |
        +---------+---------+
        v                   v
    models.py           services.py
    Product/Category    Strategy descuento
        |
        v
      MySQL
        """
    ).strip("\n")
    font = load_font(17)
    lines = diagram.splitlines()
    line_h = 26
    width, height = 720, 40 + len(lines) * line_h
    img = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)
    draw.text((12, 10), "Diagrama — Arquitectura en capas", fill="#1a1a1a", font=load_font(18))
    for i, line in enumerate(lines):
        draw.text((24, 44 + i * line_h), line, fill="#222222", font=font)
    out = CAPTURAS / "captura_diagrama_arquitectura.png"
    img.save(out)
    return out


def ensure_demo_data() -> tuple[int, int]:
    os.environ.setdefault("DJANGO_SETTINGS_MODULE", "tienda_online.settings")
    import django

    django.setup()
    from django.contrib.auth.models import User
    from products.models import Category, Product

    user, created = User.objects.get_or_create(username="demo_api", defaults={})
    if created:
        user.set_password("demo1234")
        user.save()
    else:
        user.set_password("demo1234")
        user.save()

    cat_e, _ = Category.objects.get_or_create(name="Electronica")
    cat_r, _ = Category.objects.get_or_create(name="Ropa")

    prod_e, _ = Product.objects.update_or_create(
        name="Laptop Captura",
        defaults={
            "category": cat_e,
            "description": "Equipo electronica",
            "price": "10000.00",
            "is_active": True,
        },
    )
    prod_r, _ = Product.objects.update_or_create(
        name="Camisa Captura",
        defaults={
            "category": cat_r,
            "description": "Prenda ropa",
            "price": "10000.00",
            "is_active": True,
        },
    )
    return prod_e.id, prod_r.id


def fetch_api_captures(id_e: int, id_r: int) -> tuple[Path, Path]:
    base = "http://127.0.0.1:8000"
    token_resp = requests.post(
        f"{base}/api/token/",
        json={"username": "demo_api", "password": "demo1234"},
        timeout=15,
    )
    token_resp.raise_for_status()
    access = token_resp.json()["access"]
    headers = {"Authorization": f"Bearer {access}"}

    r1 = requests.get(f"{base}/api/v1/products/{id_e}/precio-final/", headers=headers, timeout=15)
    r2 = requests.get(f"{base}/api/v1/products/{id_r}/precio-final/", headers=headers, timeout=15)
    r1.raise_for_status()
    r2.raise_for_status()

    img1 = render_json_api_image(
        "GET",
        f"{base}/api/v1/products/{id_e}/precio-final/",
        r1.status_code,
        r1.json(),
        "Authorization: Bearer <token JWT>",
    )
    img2 = render_json_api_image(
        "GET",
        f"{base}/api/v1/products/{id_r}/precio-final/",
        r2.status_code,
        r2.json(),
        "Authorization: Bearer <token JWT>",
    )
    return img1, img2


def run_tests_capture() -> Path:
    proc = subprocess.run(
        [sys.executable, "manage.py", "test", "products"],
        cwd=BASE,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
    )
    combined = (proc.stdout or "") + (proc.stderr or "")
    # Filtrar líneas útiles para la captura
    keep = []
    for line in combined.splitlines():
        if any(
            k in line
            for k in (
                "manage.py test",
                "Creating test",
                "Ran ",
                "OK",
                "FAIL",
                "Found ",
                "System check",
                "ERROR",
                "------",
            )
        ):
            keep.append(line.strip())
    if not keep:
        keep = combined.splitlines()[-15:]
    return render_terminal_image("\n".join(keep))


def add_picture(doc, img_path: Path, width_in=6.2):
    doc.add_picture(str(img_path), width=Inches(width_in))
    doc.add_paragraph()


def build_word(images: dict[str, Path]):
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    h = doc.add_heading(
        "Resultado de Aprendizaje — Unidad 1: Definición del proceso de desarrollo WEB",
        0,
    )
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(
        "Arquitectura en capas y patrones de diseño — proyecto Django REST «tienda_online»."
    )

    doc.add_heading("1. Modelo de arquitectura del proyecto", level=1)
    doc.add_paragraph(
        "El proyecto utiliza MVT (Django) con API REST (DRF) y una capa de servicios "
        "para la lógica de negocio. Flujo: urls.py → views.py → services.py → models.py → MySQL."
    )
    doc.add_heading("1.1 Diagrama del flujo", level=2)
    add_picture(doc, images["diagrama"])

    doc.add_heading("1.2 Configuración relevante", level=2)
    t = doc.add_table(rows=1, cols=2)
    t.style = "Table Grid"
    t.rows[0].cells[0].text = "Parámetro"
    t.rows[0].cells[1].text = "Valor"
    for c in t.rows[0].cells:
        set_cell_shading(c, "D9E2F3")
    for a, b in [
        ("Proyecto", "tienda_online"),
        ("Base de datos", "MySQL — django_api"),
        ("API", "/api/v1/"),
        ("JWT", "POST /api/token/"),
        ("Endpoint Facade", "GET /api/v1/products/<pk>/precio-final/"),
    ]:
        row = t.add_row().cells
        row[0].text = a
        row[1].text = b
    doc.add_paragraph()

    doc.add_heading("2. Capa de servicios — Strategy", level=1)
    add_picture(doc, images["services"])

    doc.add_heading("3. Factory Method en models.py", level=1)
    add_picture(doc, images["models"])

    doc.add_heading("4. Facade — endpoint /precio-final/", level=1)
    doc.add_paragraph(
        "Prueba con dos productos de categorías distintas (descuento 10 % vs 20 %):"
    )
    add_picture(doc, images["api_e"])
    add_picture(doc, images["api_r"])

    doc.add_heading("5. Pruebas unitarias", level=1)
    add_picture(doc, images["tests"])
    add_picture(doc, images["terminal"])

    doc.add_heading("5.1 Explicación pruebas /precio-final/", level=2)
    tt = doc.add_table(rows=1, cols=2)
    tt.style = "Table Grid"
    tt.rows[0].cells[0].text = "Prueba"
    tt.rows[0].cells[1].text = "Descripción"
    for c in tt.rows[0].cells:
        set_cell_shading(c, "D9E2F3")
    for n, d in [
        (
            "test_precio_final_autenticado",
            "Usuario autenticado recibe 200 y el JSON incluye precio_final.",
        ),
        (
            "test_precio_final_sin_autenticar",
            "Sin token responde 401 por IsAuthenticated.",
        ),
    ]:
        r = tt.add_row().cells
        r[0].text = n
        r[1].text = d

    doc.add_heading("6. Tabla de patrones de diseño", level=1)
    pat = doc.add_table(rows=1, cols=4)
    pat.style = "Table Grid"
    headers = ["Patrón", "Categoría", "Evidencia", "Explicación"]
    for i, htxt in enumerate(headers):
        pat.rows[0].cells[i].text = htxt
        set_cell_shading(pat.rows[0].cells[i], "D9E2F3")
    rows = [
        (
            "Factory Method",
            "Creacional",
            "Product.crear_con_categoria_default()",
            "Centraliza creación con categoría «Sin categoría» mediante get_or_create.",
        ),
        (
            "Strategy",
            "Comportamiento",
            "services.py — estrategias de descuento",
            "Algoritmo de descuento intercambiable según nombre de categoría.",
        ),
        (
            "Facade",
            "Estructural",
            "views.precio_final + ruta precio-final",
            "Un solo GET oculta modelo, estrategia y cálculo del precio final.",
        ),
    ]
    for row_data in rows:
        row = pat.add_row().cells
        for i, val in enumerate(row_data):
            row[i].text = val

    doc.save(OUTPUT)


def main():
    CAPTURAS.mkdir(exist_ok=True)

    images = {
        "diagrama": render_diagram_image(),
        "services": render_code_image(
            BASE / "products" / "services.py",
            title="products/services.py",
        ),
        "models": render_code_image(
            BASE / "products" / "models.py",
            title="products/models.py — Factory Method",
            start_line=35,
            end_line=46,
            highlight_lines={38, 39, 40, 41, 42, 43, 44, 45, 46},
        ),
        "tests": render_code_image(
            BASE / "products" / "tests.py",
            title="products/tests.py — pruebas precio-final",
            start_line=58,
            end_line=71,
            highlight_lines=set(range(63, 72)),
        ),
    }

    images["terminal"] = run_tests_capture()

    id_e, id_r = ensure_demo_data()
    server = subprocess.Popen(
        [sys.executable, "manage.py", "runserver", "127.0.0.1:8000", "--noreload"],
        cwd=BASE,
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )
    try:
        time.sleep(2.5)
        api_e, api_r = fetch_api_captures(id_e, id_r)
        images["api_e"] = api_e
        images["api_r"] = api_r
    finally:
        server.terminate()
        try:
            server.wait(timeout=5)
        except subprocess.TimeoutExpired:
            server.kill()

    build_word(images)
    print(f"Listo: {OUTPUT}")
    print("Capturas en:", CAPTURAS)


if __name__ == "__main__":
    main()
