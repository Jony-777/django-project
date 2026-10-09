# -*- coding: utf-8 -*-
"""Genera el documento Word del Resultado de Aprendizaje Unidad 1."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor, Inches
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUTPUT = Path(__file__).resolve().parent / "Resultado_Aprendizaje_Unidad1.docx"


def set_cell_shading(cell, fill_hex: str):
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill_hex)
    cell._tc.get_or_add_tcPr().append(shading)


def add_placeholder(doc, text: str):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.italic = True
    run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    p.paragraph_format.space_after = Pt(12)


def add_code_block(doc, code: str):
    for line in code.splitlines():
        p = doc.add_paragraph(line)
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_after = Pt(0)
        for run in p.runs:
            run.font.name = "Consolas"
            run.font.size = Pt(9)
    doc.add_paragraph()


def main():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    title = doc.add_heading(
        "Resultado de Aprendizaje — Unidad 1: Definición del proceso de desarrollo WEB",
        level=0,
    )
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(
        "Arquitectura en capas y patrones de diseño en el proyecto Django REST "
        "«tienda_online» (Prácticas 2 a 5 + capa de servicios)."
    )
    doc.add_paragraph()

    # --- 1 ---
    doc.add_heading("1. Modelo de arquitectura del proyecto", level=1)
    doc.add_paragraph(
        "El proyecto sigue el patrón MVT (Model–View–Template) propio de Django, "
        "adaptado a una API REST: no hay plantillas HTML; las «vistas» son controladores "
        "de Django REST Framework que devuelven JSON. Sobre ese esquema se añadió una "
        "arquitectura en capas explícita, separando la lógica de negocio del acceso HTTP."
    )
    doc.add_paragraph(
        "Capas identificadas:",
        style="List Bullet",
    )
    layers = [
        "Enrutamiento: tienda_online/urls.py y products/urls.py (asignan URL → vista).",
        "Presentación / controlador: products/views.py (ViewSets y vista precio_final; JWT).",
        "Lógica de negocio: products/services.py (estrategias de descuento).",
        "Persistencia: products/models.py (ORM) y MySQL (base django_api).",
        "Serialización: products/serializers.py (entrada/salida JSON).",
    ]
    for item in layers:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("1.1 Diagrama del flujo de una petición (precio final)", level=2)
    doc.add_paragraph(
        "Flujo para GET /api/v1/products/<id>/precio-final/ con token JWT:"
    )
    diagram = """
    ┌─────────────┐     ┌──────────────────┐     ┌─────────────────┐
    │   Cliente   │────▶│  urls.py         │────▶│  views.py       │
    │  (Postman)  │     │  (enrutamiento)  │     │  precio_final   │
    └─────────────┘     └──────────────────┘     └────────┬────────┘
                                                          │
                        ┌─────────────────┐               │
                        │  models.py      │◀──────────────┤
                        │  Product,       │               │
                        │  Category       │               ▼
                        └────────┬────────┘     ┌─────────────────┐
                                 │              │  services.py    │
                                 ▼              │  Strategy:      │
                        ┌─────────────────┐     │  obtener_       │
                        │     MySQL       │     │  estrategia()   │
                        └─────────────────┘     └─────────────────┘
    """
    add_code_block(doc, diagram.strip())
    add_placeholder(
        doc,
        "[INSERTAR CAPTURA OPCIONAL: diagrama elaborado en draw.io, PowerPoint o Canva, "
        "equivalente al flujo anterior.]",
    )

    doc.add_heading("1.2 Configuración relevante (reporte breve)", level=2)
    table_cfg = doc.add_table(rows=1, cols=2)
    table_cfg.style = "Table Grid"
    hdr = table_cfg.rows[0].cells
    hdr[0].text = "Parámetro"
    hdr[1].text = "Valor"
    for c in hdr:
        set_cell_shading(c, "D9E2F3")
    rows_cfg = [
        ("Proyecto Django", "tienda_online"),
        ("App de dominio", "products"),
        ("Motor de BD", "MySQL (django.db.backends.mysql)"),
        ("Base de datos", "django_api"),
        ("Prefijo API REST", "/api/v1/"),
        ("Autenticación API", "JWT (djangorestframework-simplejwt)"),
        ("Token", "POST /api/token/"),
        ("Endpoint nuevo", "GET /api/v1/products/<pk>/precio-final/"),
    ]
    for a, b in rows_cfg:
        row = table_cfg.add_row().cells
        row[0].text = a
        row[1].text = b
    doc.add_paragraph()

    # --- 2 services ---
    doc.add_heading("2. Capa de servicios y patrón Strategy (comportamiento)", level=1)
    doc.add_paragraph(
        "Archivo products/services.py: define la interfaz EstrategiaDescuento y "
        "implementaciones según categoría (Electrónica 10 %, Ropa 20 %, resto sin descuento). "
        "La función obtener_estrategia(categoria_nombre) selecciona la estrategia; "
        "el nombre de categoría debe coincidir exactamente con el guardado en la BD "
        "(p. ej. «Electronica», «Ropa»)."
    )
    add_placeholder(doc, "[INSERTAR CAPTURA: archivo products/services.py en el IDE.]")
    doc.add_heading("Referencia de código (services.py)", level=3)
    add_code_block(
        doc,
        Path(__file__).parent.joinpath("products/services.py").read_text(encoding="utf-8"),
    )

    # --- 3 Factory ---
    doc.add_heading("3. Patrón Factory Method (creacional)", level=1)
    doc.add_paragraph(
        "En products/models.py, el método de clase Product.crear_con_categoria_default "
        "centraliza la creación de productos cuando no se indica categoría: usa "
        "get_or_create sobre la categoría «Sin categoría» y luego create(). "
        "Corresponde a Factory Method porque delega la decisión de qué «tipo» de "
        "producto (con categoría por defecto) instanciar en un único punto del modelo."
    )
    add_placeholder(
        doc,
        "[INSERTAR CAPTURA: método crear_con_categoria_default en products/models.py.]",
    )
    doc.add_heading("Referencia de código (Factory Method)", level=3)
    add_code_block(
        doc,
        """    @classmethod
    def crear_con_categoria_default(cls, name, description, price):
        categoria_default, _ = Category.objects.get_or_create(name='Sin categoría')
        return cls.objects.create(
            name=name,
            description=description,
            price=price,
            category=categoria_default,
        )""",
    )

    # --- 4 Facade ---
    doc.add_heading("4. Patrón Facade (estructural) y endpoint /precio-final/", level=1)
    doc.add_paragraph(
        "La vista precio_final en products/views.py actúa como fachada: el cliente "
        "solo consume un GET con el id del producto; internamente se obtiene el producto, "
        "se elige la estrategia de descuento y se calcula el precio final. "
        "El endpoint está protegido con IsAuthenticated (JWT), igual que en la Práctica 4."
    )
    doc.add_paragraph("Ruta registrada en products/urls.py:")
    add_code_block(
        doc,
        "path('products/<int:pk>/precio-final/', precio_final),",
    )
    doc.add_heading("Referencia de código (Facade — vista precio_final)", level=3)
    add_code_block(
        doc,
        """@api_view(['GET'])
@permission_classes([IsAuthenticated])
def precio_final(request, pk):
    producto = Product.objects.get(pk=pk)
    estrategia = obtener_estrategia(producto.category.name)
    precio_con_descuento = estrategia.calcular(float(producto.price))
    return Response({
        'producto': producto.name,
        'precio_original': producto.price,
        'precio_final': round(precio_con_descuento, 2),
    })""",
    )

    doc.add_heading("4.1 Prueba manual del endpoint (mínimo 2 productos)", level=2)
    doc.add_paragraph("Pasos recomendados para Postman o cliente HTTP:")
    steps = [
        "Iniciar servidor: python manage.py runserver",
        "Obtener token: POST http://127.0.0.1:8000/api/token/ con body JSON "
        '{"username": "...", "password": "..."}',
        "Header en las peticiones GET: Authorization: Bearer <access_token>",
        "Producto categoría Electronica: GET http://127.0.0.1:8000/api/v1/products/<id>/precio-final/",
        "Producto categoría Ropa: misma URL con otro id; precio_final debe ser menor (20 % vs 10 %).",
    ]
    for s in steps:
        doc.add_paragraph(s, style="List Number")

    doc.add_paragraph("Ejemplo de respuesta esperada (Electrónica, precio 10000):")
    add_code_block(
        doc,
        '{"producto": "Laptop", "precio_original": "10000.00", "precio_final": 9000.0}',
    )
    doc.add_paragraph("Ejemplo (Ropa, precio 10000):")
    add_code_block(
        doc,
        '{"producto": "Camisa", "precio_original": "10000.00", "precio_final": 8000.0}',
    )
    add_placeholder(
        doc,
        "[INSERTAR CAPTURA 1: respuesta JSON producto categoría Electronica.]",
    )
    add_placeholder(
        doc,
        "[INSERTAR CAPTURA 2: respuesta JSON producto categoría Ropa (descuento distinto).]",
    )

    # --- 5 Tests ---
    doc.add_heading("5. Pruebas unitarias (Práctica 5 + endpoint precio-final)", level=1)
    doc.add_paragraph(
        "En products/tests.py se mantienen las 6 pruebas de la Unidad 1 (CRUD categorías "
        "y productos) y se agregaron 2 pruebas para el nuevo endpoint. Total: 8 pruebas."
    )
    add_placeholder(doc, "[INSERTAR CAPTURA: tests.py mostrando test_precio_final_autenticado y test_precio_final_sin_autenticar.]")
    add_placeholder(
        doc,
        "[INSERTAR CAPTURA: terminal con python manage.py test — debe mostrar «Ran 8 tests» y «OK».]",
    )

    doc.add_heading("5.1 Explicación de las pruebas del endpoint", level=2)
    test_table = doc.add_table(rows=1, cols=2)
    test_table.style = "Table Grid"
    h = test_table.rows[0].cells
    h[0].text = "Prueba"
    h[1].text = "Qué verifica"
    for c in h:
        set_cell_shading(c, "D9E2F3")
    tests_explain = [
        (
            "test_precio_final_autenticado",
            "Con usuario autenticado (force_authenticate en setUp), GET a "
            "/api/v1/products/<id>/precio-final/ debe responder 200 y el JSON debe "
            "incluir la clave precio_final.",
        ),
        (
            "test_precio_final_sin_autenticar",
            "Tras force_authenticate(user=None), la misma URL debe responder 401 "
            "porque el endpoint exige JWT.",
        ),
    ]
    for name, desc in tests_explain:
        r = test_table.add_row().cells
        r[0].text = name
        r[1].text = desc

    doc.add_paragraph()
    doc.add_heading("5.2 Resumen de las 6 pruebas previas (Práctica 5)", level=2)
    prev = doc.add_table(rows=1, cols=2)
    prev.style = "Table Grid"
    ph = prev.rows[0].cells
    ph[0].text = "Prueba"
    ph[1].text = "Qué verifica"
    for c in ph:
        set_cell_shading(c, "D9E2F3")
    for name, desc in [
        ("test_create_category", "POST crea una categoría nueva (201)."),
        ("test_list_categories", "GET lista categorías existentes (200)."),
        ("test_create_product", "POST crea un producto con category_id (201)."),
        ("test_list_products", "GET lista productos (200)."),
        ("test_update_product", "PUT actualiza un producto (200)."),
        ("test_delete_product", "DELETE elimina un producto (204)."),
    ]:
        row = prev.add_row().cells
        row[0].text = name
        row[1].text = desc

    doc.add_heading("Referencia de código (pruebas nuevas)", level=3)
    add_code_block(
        doc,
        """    def test_precio_final_autenticado(self):
        response = self.client.get(f'/api/v1/products/{self.product.id}/precio-final/')
        self.assertEqual(response.status_code, 200)
        self.assertIn('precio_final', response.data)

    def test_precio_final_sin_autenticar(self):
        self.client.force_authenticate(user=None)
        response = self.client.get(f'/api/v1/products/{self.product.id}/precio-final/')
        self.assertEqual(response.status_code, 401)""",
    )

    # --- 6 Patterns table ---
    doc.add_heading("6. Tabla de patrones de diseño implementados", level=1)
    pat = doc.add_table(rows=1, cols=4)
    pat.style = "Table Grid"
    hp = pat.rows[0].cells
    for i, t in enumerate(["Patrón", "Categoría", "Evidencia (archivo / elemento)", "Explicación"]):
        hp[i].text = t
        set_cell_shading(hp[i], "D9E2F3")
    patterns = [
        (
            "Factory Method",
            "Creacional",
            "products/models.py — Product.crear_con_categoria_default()",
            "Encapsula la creación de productos sin categoría explícita: "
            "obtiene o crea la categoría «Sin categoría» y devuelve una instancia de Product. "
            "Evita duplicar lógica de construcción en vistas o scripts.",
        ),
        (
            "Strategy",
            "Comportamiento",
            "products/services.py — EstrategiaDescuento, DescuentoElectronica, "
            "DescuentoRopa, SinDescuento, obtener_estrategia()",
            "Intercambia el algoritmo de descuento según el nombre de la categoría "
            "sin condicionales en la vista. Cada estrategia implementa calcular(precio); "
            "la vista solo delega en la estrategia elegida.",
        ),
        (
            "Facade",
            "Estructural",
            "products/views.py — función precio_final; ruta precio-final en urls.py",
            "Ofrece una interfaz simple (un GET, un JSON) que oculta la coordinación "
            "entre modelo, selección de estrategia y cálculo numérico. "
            "El cliente no necesita conocer services.py ni las clases de descuento.",
        ),
    ]
    for row_data in patterns:
        row = pat.add_row().cells
        for i, val in enumerate(row_data):
            row[i].text = val

    doc.add_paragraph()
    doc.add_heading("7. Integración con la Unidad 1", level=1)
    doc.add_paragraph(
        "Este resultado de aprendizaje cierra el ciclo de la unidad: metodología ágil "
        "(organización del trabajo en prácticas incrementales), desarrollo con Django REST "
        "y MySQL (Prácticas 2–3), seguridad con JWT (Práctica 4), calidad con pruebas "
        "automáticas (Práctica 5) y ahora separación en capas con patrones creacionales, "
        "de comportamiento y estructurales aplicados al mismo proyecto tienda_online."
    )

    doc.add_paragraph()
    p_end = doc.add_paragraph()
    p_end.add_run(
        "Nota: Sustituya cada texto [INSERTAR CAPTURA...] por la imagen correspondiente "
        "antes de exportar a PDF. La portada la elabora el estudiante por separado."
    ).italic = True

    doc.save(OUTPUT)
    print(f"Documento generado: {OUTPUT}")


if __name__ == "__main__":
    main()
